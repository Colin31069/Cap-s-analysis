from __future__ import annotations

from collections import OrderedDict
from concurrent.futures import CancelledError
import os
from typing import Callable

import numpy as np

from .config import (
    BASELINE_RISE_MIN_RUN,
    BASELINE_RISE_ROLLING_POINTS,
    BASELINE_TAIL_POINTS,
    DATA_COL,
    DEFAULT_BASELINE_DURATION_SEC,
    DEFAULT_BASELINE_WARNING_THRESHOLD_PCT,
    DEFAULT_DRUG_APPLY_TIME_SEC,
    DEFAULT_DRUG_APPLY_TOLERANCE_SEC,
    DROP_SIGMA_THRESHOLD,
    DT_SEC,
    LEGACY_DROP_SEARCH_POINTS,
)
from .models import BaselineWarningStatus, DropDetectionSource, ProcessedSignal


def check_cancelled(cancelled: Callable[[], bool] | None) -> None:
    if cancelled is not None and cancelled():
        raise CancelledError("分析已取消。")


class CapacitanceCache:
    """Bounded in-memory cache; use from the single analysis worker only."""

    def __init__(self, max_bytes: int = 64 * 1024 * 1024) -> None:
        self.max_bytes = max_bytes
        self._entries: OrderedDict[tuple, np.ndarray] = OrderedDict()
        self._bytes = 0

    def load(self, path: str, cancelled: Callable[[], bool] | None = None) -> np.ndarray:
        check_cancelled(cancelled)
        before = os.stat(path)
        key = (os.path.abspath(path), before.st_size, before.st_mtime_ns, before.st_ctime_ns)
        if key in self._entries:
            self._entries.move_to_end(key)
            return self._entries[key]
        values = load_capacitance_values(path, cancelled=cancelled)
        after = os.stat(path)
        if (after.st_size, after.st_mtime_ns) != key[1:3]:
            raise ValueError("讀取期間檔案已變更，請等待同步完成後重新製圖。")
        if after.st_ctime_ns != before.st_ctime_ns:
            # Cloud hydration can update metadata without changing the content.
            # Use this read, but require a stable read before caching it.
            return values
        if values.nbytes <= self.max_bytes:
            for old in [entry for entry in self._entries if entry[0] == key[0]]:
                self._bytes -= self._entries.pop(old).nbytes
            while self._entries and self._bytes + values.nbytes > self.max_bytes:
                self._bytes -= self._entries.popitem(last=False)[1].nbytes
            values.setflags(write=False)
            self._entries[key] = values
            self._bytes += values.nbytes
        return values


def load_capacitance_values(path: str, *, cancelled: Callable[[], bool] | None = None) -> np.ndarray:
    # Excel support is loaded only when a workbook is actually opened.
    from openpyxl import load_workbook

    check_cancelled(cancelled)
    with open(path, "rb") as file:
        workbook = load_workbook(file, read_only=True, data_only=True)
        try:
            sheet = workbook.worksheets[0]
            header = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
            if DATA_COL not in header:
                raise ValueError(f"缺少欄位：{DATA_COL}")
            column = header.index(DATA_COL) + 1
            # Follow actual worksheet rows, not an inflated saved used-range.
            sheet.reset_dimensions()
            values = []
            for index, row in enumerate(sheet.iter_rows(min_row=2, min_col=column, max_col=column, values_only=True)):
                if index % 128 == 0:
                    check_cancelled(cancelled)
                values.append(row[0])
            while values and values[-1] is None:
                values.pop()
            if not values:
                raise ValueError("沒有可分析的資料。")
            numeric: list[float] = []
            for row, value in enumerate(values, 2):
                if row % 128 == 0:
                    check_cancelled(cancelled)
                if value is None or isinstance(value, bool):
                    raise ValueError(f"第 {row} 列不是數值；請修正資料，避免改變時間軸。")
                try:
                    number = float(value)
                except (TypeError, ValueError) as exc:
                    raise ValueError(f"第 {row} 列不是數值。") from exc
                if not np.isfinite(number):
                    raise ValueError(f"第 {row} 列包含非有限數值。")
                numeric.append(number)
            return np.asarray(numeric, dtype=float)
        finally:
            workbook.close()


def read_xlsx_single(path: str) -> np.ndarray | None:
    try:
        return load_capacitance_values(path)
    except Exception:
        return None


def _offset_pct(value: float, baseline: float) -> float:
    if abs(baseline) < 1e-12:
        return 0.0
    return ((value / baseline) * 100.0) - 100.0


def _seconds_to_points(duration_sec: float, minimum: int = 0, sample_interval_sec: float = DT_SEC) -> int:
    return max(minimum, int(round(duration_sec / sample_interval_sec)))


def _first_threshold_crossing(
    samples: np.ndarray,
    threshold: float,
    start_idx: int,
    end_idx: int,
) -> int | None:
    if start_idx >= end_idx:
        return None

    search_range = samples[start_idx:end_idx]
    over_idx = np.where(search_range > threshold)[0]
    if over_idx.size == 0:
        return None
    return start_idx + int(over_idx[0])


def _max_continuous_rise_offset_pct(baseline_segment: np.ndarray, baseline_avg: float) -> float:
    if baseline_segment.size < BASELINE_RISE_ROLLING_POINTS:
        return 0.0

    kernel = np.ones(BASELINE_RISE_ROLLING_POINTS, dtype=float) / BASELINE_RISE_ROLLING_POINTS
    rolling_mean = np.convolve(baseline_segment, kernel, mode="valid")
    if rolling_mean.size < BASELINE_RISE_MIN_RUN:
        return 0.0

    run_length = 1
    max_offset_pct = 0.0

    for index in range(1, len(rolling_mean)):
        if rolling_mean[index] > rolling_mean[index - 1]:
            run_length += 1
            continue

        if run_length >= BASELINE_RISE_MIN_RUN:
            max_offset_pct = max(max_offset_pct, _offset_pct(float(rolling_mean[index - 1]), baseline_avg))
        run_length = 1

    if run_length >= BASELINE_RISE_MIN_RUN:
        max_offset_pct = max(max_offset_pct, _offset_pct(float(rolling_mean[-1]), baseline_avg))

    return max(0.0, max_offset_pct)


def _baseline_warning_status(tail_hit: bool, rise_hit: bool) -> BaselineWarningStatus:
    hit_count = int(tail_hit) + int(rise_hit)
    if hit_count >= 2:
        return "inaccurate"
    if hit_count == 1:
        return "warning"
    return "ok"


def _effective_baseline_points(
    sample_count: int,
    requested_baseline_duration_sec: float,
    apply_window_start_idx: int,
    sample_interval_sec: float = DT_SEC,
) -> tuple[int, bool]:
    requested_baseline_points = min(
        sample_count,
        _seconds_to_points(max(0.0, requested_baseline_duration_sec), minimum=1, sample_interval_sec=sample_interval_sec),
    )
    if requested_baseline_points <= apply_window_start_idx:
        return requested_baseline_points, False

    shortened_points = min(sample_count, apply_window_start_idx)
    if shortened_points <= 0:
        return min(sample_count, 1), True
    return shortened_points, True


def _build_timing_warning_details(
    requested_baseline_duration_sec: float,
    effective_baseline_duration_sec: float,
    baseline_was_auto_shortened: bool,
    apply_time_sec: float,
    apply_tolerance_sec: float,
    drop_search_fallback_used: bool,
    sample_interval_sec: float = DT_SEC,
) -> tuple[str, ...]:
    details: list[str] = []

    if baseline_was_auto_shortened:
        if effective_baseline_duration_sec <= sample_interval_sec:
            details.append(
                "Baseline overlapped the apply window and was clamped to the minimum baseline length."
            )
        else:
            details.append(
                f"Baseline shortened from {requested_baseline_duration_sec:.1f}s to "
                f"{effective_baseline_duration_sec:.1f}s to stay before the apply window."
            )

    if drop_search_fallback_used:
        details.append(
            f"No threshold crossing found within {apply_time_sec:.1f}s +/- {apply_tolerance_sec:.1f}s; "
            "used automatic fallback search."
        )

    return tuple(details)


def _fallback_drop_index(samples: np.ndarray, threshold: float, baseline_end_idx: int) -> int:
    search_end = min(len(samples), baseline_end_idx + LEGACY_DROP_SEARCH_POINTS)
    crossing_idx = _first_threshold_crossing(samples, threshold, baseline_end_idx, search_end)
    if crossing_idx is None:
        return baseline_end_idx
    return crossing_idx


def analyze_signal(
    data: np.ndarray,
    baseline_warning_threshold_pct: float = DEFAULT_BASELINE_WARNING_THRESHOLD_PCT,
    baseline_duration_sec: float = DEFAULT_BASELINE_DURATION_SEC,
    drug_apply_time_sec: float = DEFAULT_DRUG_APPLY_TIME_SEC,
    drug_apply_tolerance_sec: float = DEFAULT_DRUG_APPLY_TOLERANCE_SEC,
    sample_interval_sec: float = DT_SEC,
) -> ProcessedSignal | None:
    samples = np.asarray(data, dtype=float)
    if samples.size == 0:
        return None
    if not np.isfinite(sample_interval_sec) or sample_interval_sec <= 0:
        raise ValueError("取樣間隔必須是大於 0 的有限數值。")

    warning_threshold_pct = max(0.0, float(baseline_warning_threshold_pct))
    requested_baseline_duration_sec = max(float(baseline_duration_sec), sample_interval_sec)
    requested_apply_time_sec = max(0.0, float(drug_apply_time_sec))
    requested_apply_tolerance_sec = max(0.0, float(drug_apply_tolerance_sec))

    apply_center_idx = _seconds_to_points(requested_apply_time_sec, sample_interval_sec=sample_interval_sec)
    apply_tolerance_points = _seconds_to_points(requested_apply_tolerance_sec, sample_interval_sec=sample_interval_sec)
    apply_window_start_idx = max(0, apply_center_idx - apply_tolerance_points)
    apply_window_end_idx = min(len(samples), apply_center_idx + apply_tolerance_points + 1)

    baseline_len, baseline_was_auto_shortened = _effective_baseline_points(
        len(samples),
        requested_baseline_duration_sec,
        apply_window_start_idx,
        sample_interval_sec,
    )
    baseline_segment = samples[:baseline_len]
    initial_avg = float(np.mean(baseline_segment))
    sample_std = float(np.std(baseline_segment, ddof=1)) if baseline_len > 1 else 0.0
    effective_baseline_duration_sec = baseline_len * sample_interval_sec

    tail_len = min(len(baseline_segment), BASELINE_TAIL_POINTS)
    tail_mean = float(np.mean(baseline_segment[-tail_len:])) if tail_len > 0 else initial_avg
    tail_offset_pct = _offset_pct(tail_mean, initial_avg)
    rise_offset_pct = _max_continuous_rise_offset_pct(baseline_segment, initial_avg)
    tail_warning_hit = abs(tail_offset_pct) >= warning_threshold_pct
    rise_warning_hit = rise_offset_pct >= warning_threshold_pct
    warning_status = _baseline_warning_status(tail_warning_hit, rise_warning_hit)

    threshold = initial_avg + DROP_SIGMA_THRESHOLD * max(sample_std, 1e-4)
    window_search_start_idx = max(baseline_len, apply_window_start_idx)
    idx_drop = _first_threshold_crossing(
        samples,
        threshold,
        window_search_start_idx,
        apply_window_end_idx,
    )
    drop_detection_source: DropDetectionSource = "window"
    drop_search_fallback_used = idx_drop is None

    if idx_drop is None:
        idx_drop = _fallback_drop_index(samples, threshold, baseline_len)
        drop_detection_source = "fallback_auto"

    final_avg = float(np.nanmean(samples[-100:])) if len(samples) >= 100 else float("nan")
    delta_raw = final_avg - initial_avg if not np.isnan(final_avg) else float("nan")
    timing_warning_details = _build_timing_warning_details(
        requested_baseline_duration_sec=requested_baseline_duration_sec,
        effective_baseline_duration_sec=effective_baseline_duration_sec,
        baseline_was_auto_shortened=baseline_was_auto_shortened,
        apply_time_sec=requested_apply_time_sec,
        apply_tolerance_sec=requested_apply_tolerance_sec,
        drop_search_fallback_used=drop_search_fallback_used,
        sample_interval_sec=sample_interval_sec,
    )

    return ProcessedSignal(
        time_sec=np.arange(len(samples), dtype=float) * sample_interval_sec,
        capacitance=samples,
        drop_time=idx_drop * sample_interval_sec,
        delta_capacitance=delta_raw,
        initial_avg=initial_avg,
        effective_baseline_points=baseline_len,
        effective_baseline_duration_sec=effective_baseline_duration_sec,
        baseline_was_auto_shortened=baseline_was_auto_shortened,
        drop_detection_source=drop_detection_source,
        drop_search_fallback_used=drop_search_fallback_used,
        timing_warning_details=timing_warning_details,
        baseline_warning_status=warning_status,
        baseline_tail_offset_pct=tail_offset_pct,
        baseline_rise_offset_pct=rise_offset_pct,
        baseline_tail_warning_hit=tail_warning_hit,
        baseline_rise_warning_hit=rise_warning_hit,
    )


def process_single_file(
    file_path: str,
    baseline_warning_threshold_pct: float = DEFAULT_BASELINE_WARNING_THRESHOLD_PCT,
    baseline_duration_sec: float = DEFAULT_BASELINE_DURATION_SEC,
    drug_apply_time_sec: float = DEFAULT_DRUG_APPLY_TIME_SEC,
    drug_apply_tolerance_sec: float = DEFAULT_DRUG_APPLY_TOLERANCE_SEC,
    sample_interval_sec: float = DT_SEC,
) -> ProcessedSignal | None:
    samples = read_xlsx_single(file_path)
    if samples is None or samples.size == 0:
        return None
    return analyze_signal(
        samples,
        baseline_warning_threshold_pct=baseline_warning_threshold_pct,
        baseline_duration_sec=baseline_duration_sec,
        drug_apply_time_sec=drug_apply_time_sec,
        drug_apply_tolerance_sec=drug_apply_tolerance_sec,
        sample_interval_sec=sample_interval_sec,
    )
